"""
excel_generator.py — Genera .xlsx y .csv en el formato exacto de la organización.

Formato calibrado desde plantilla: Formato Deck Pruebas-JIRA.xlsx
Columnas: Name | Objective | Precondition | Estimated Time | Step | Data | Expected Result
- Una fila por caso (NO una fila por paso)
- Step: todos los pasos numerados en una sola celda con saltos de línea
- Estimated Time: formato HH:MM
- CSV derivado EXCLUSIVAMENTE del Excel generado (nunca desde memoria)
"""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any, Optional

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"],
                          stdout=subprocess.DEVNULL)
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter


# ----------------------------------------------------------------- estilos
HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT  = Font(color="FFFFFF", bold=True, size=10, name="Calibri")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

THIN = Side(style="thin")
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DATA_ALIGN  = Alignment(vertical="top", wrap_text=True)
DATA_FONT   = Font(size=9, name="Calibri")

# Columnas en el orden exacto de la plantilla
COLUMNS = [
    "Name",
    "Objective",
    "Precondition",
    "Estimated Time",
    "Step",
    "Data",
    "Expected Result",
]

# Anchos de columna calibrados para legibilidad
COL_WIDTHS = [45, 55, 55, 12, 60, 40, 60]


def _fmt_time(minutes: Any) -> str:
    """Convierte minutos a HH:MM."""
    try:
        m = int(minutes)
        return f"{m // 60:02d}:{m % 60:02d}"
    except (TypeError, ValueError):
        return "00:30"


def _fmt_steps(steps: list[str]) -> str:
    """Numera los pasos dentro de una sola celda con saltos de línea."""
    if not steps:
        return ""
    return "\n".join(f"{i}. {s.strip()}" for i, s in enumerate(steps, 1))


def _fmt_precondition(precondition: str) -> str:
    """Formatea la precondición como lista con * bullets si tiene múltiples líneas."""
    if not precondition:
        return ""
    lines = [l.strip() for l in precondition.splitlines() if l.strip()]
    if len(lines) == 1:
        return lines[0]
    return "\n".join(f"* {l}" if not l.startswith("*") else l for l in lines)


def _build_name(project: str, rule_ref: str, scenario: str, seq: int) -> str:
    """
    Construye el Name en la nomenclatura de la organización:
    PROYECTO_RuleRef_Escenario_NNN
    """
    safe_proj    = re.sub(r"[^\w]", "_", project).upper()
    safe_ref     = re.sub(r"[^\w]", "_", rule_ref).upper()
    safe_scenario= re.sub(r"[^\w\s]", "", scenario).strip()
    safe_scenario= re.sub(r"\s+", "_", safe_scenario)[:50]
    return f"{safe_proj}_{safe_ref}_{safe_scenario}_{seq:03d}"


def _apply_header(ws) -> None:
    for col, name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border    = THIN_BORDER
    ws.row_dimensions[1].height = 22


def _write_data_cell(ws, row: int, col: int, value: str) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font      = DATA_FONT
    cell.alignment = DATA_ALIGN
    cell.border    = THIN_BORDER


def _set_col_widths(ws) -> None:
    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


class ExcelGenerator:

    # ---------------------------------------------------------------- XLSX
    def generate_xlsx(
        self,
        output_path: str,
        test_cases: list[dict],
        risks: list[dict],
        traceability: list[dict],
        coverage_report,
        project: str = "PROYECTO",
    ) -> int:
        """
        Genera el Excel con el formato exacto de la organización.
        Hoja principal: Test Cases (formato Jira).
        Hojas adicionales: Risks, Traceability, Coverage.
        Retorna el número de filas de datos (sin encabezado).
        """
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        row_count = self._sheet_test_cases(wb, test_cases, project)
        self._sheet_risks(wb, risks)
        self._sheet_traceability(wb, traceability)
        self._sheet_coverage(wb, coverage_report)

        wb.save(output_path)
        return row_count

    # ---------------------------------------- Hoja Test Cases
    def _sheet_test_cases(self, wb, test_cases: list[dict], project: str) -> int:
        ws = wb.create_sheet("Test Cases")
        ws.freeze_panes = "A2"

        _apply_header(ws)
        _set_col_widths(ws)

        # Contador por regla para la nomenclatura NNN
        seq_counter: dict[str, int] = {}

        row = 2
        for tc in test_cases:
            rule_ref = tc.get("requirement_ref", "GENERAL")
            seq_counter[rule_ref] = seq_counter.get(rule_ref, 0) + 1
            seq = seq_counter[rule_ref]

            # Extraer escenario del nombre del caso (sin el prefijo de técnica)
            raw_name = tc.get("name", "")
            scenario = re.sub(r"^\[.*?\]\s*\w+-\d+\s*—\s*", "", raw_name).strip()
            if not scenario:
                scenario = raw_name

            name = _build_name(project, rule_ref, scenario, seq)

            steps_formatted      = _fmt_steps(tc.get("steps", []))
            precondition_fmt     = _fmt_precondition(tc.get("precondition", ""))
            estimated_time       = _fmt_time(tc.get("estimated_min", 30))

            row_data = [
                name,
                tc.get("objective", ""),
                precondition_fmt,
                estimated_time,
                steps_formatted,
                tc.get("test_data", ""),
                tc.get("expected_result", ""),
            ]

            for col, value in enumerate(row_data, 1):
                _write_data_cell(ws, row, col, value)

            # Altura de fila dinámica según número de pasos
            n_steps = len(tc.get("steps", []))
            ws.row_dimensions[row].height = max(40, n_steps * 15)

            row += 1

        return row - 2  # filas de datos sin encabezado

    # ---------------------------------------- Hoja Risks
    def _sheet_risks(self, wb, risks: list[dict]) -> None:
        ws = wb.create_sheet("Risks")
        ws.freeze_panes = "A2"

        headers = ["ID", "Descripcion", "Probabilidad (1-5)", "Impacto (1-5)",
                   "Score", "Nivel", "Mitigacion", "Referencia RN"]
        widths  = [8, 45, 12, 12, 8, 10, 55, 15]

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = HEADER_ALIGN
            cell.border = THIN_BORDER
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        risk_fills = {
            "ALTO":  PatternFill("solid", fgColor="FF4444"),
            "MEDIO": PatternFill("solid", fgColor="FFAA00"),
            "BAJO":  PatternFill("solid", fgColor="44BB44"),
        }
        risk_fonts = {
            "ALTO":  Font(color="FFFFFF", bold=True, size=9),
            "MEDIO": Font(color="000000", bold=True, size=9),
            "BAJO":  Font(color="FFFFFF", bold=True, size=9),
        }

        for row_idx, risk in enumerate(risks, 2):
            score = risk.get("score", 0)
            level = "ALTO" if score >= 15 else ("MEDIO" if score >= 8 else "BAJO")
            values = [
                risk.get("id", ""),
                risk.get("description", ""),
                risk.get("probability", ""),
                risk.get("impact", ""),
                score,
                level,
                risk.get("mitigation", ""),
                risk.get("rule_ref", ""),
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else "")
                cell.font      = DATA_FONT
                cell.alignment = DATA_ALIGN
                cell.border    = THIN_BORDER
                if col_idx == 6:  # columna Nivel
                    cell.fill = risk_fills.get(level, PatternFill())
                    cell.font = risk_fonts.get(level, DATA_FONT)

    # ---------------------------------------- Hoja Traceability
    def _sheet_traceability(self, wb, traceability: list[dict]) -> None:
        ws = wb.create_sheet("Traceability")
        ws.freeze_panes = "A2"

        headers = ["Regla/CA", "Descripcion", "Casos Asociados",
                   "Nivel", "Tipo", "Tecnica", "Cobertura %", "Gap"]
        widths  = [12, 45, 30, 22, 18, 22, 10, 8]

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = HEADER_ALIGN
            cell.border = THIN_BORDER
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        for row_idx, entry in enumerate(traceability, 2):
            values = [
                entry.get("rule_id", ""),
                entry.get("description", ""),
                entry.get("cases", ""),
                entry.get("level", ""),
                entry.get("type", ""),
                entry.get("technique", ""),
                entry.get("coverage_pct", ""),
                entry.get("gap", ""),
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col_idx,
                               value=str(val) if val is not None else "")
                cell.font      = DATA_FONT
                cell.alignment = DATA_ALIGN
                cell.border    = THIN_BORDER

    # ---------------------------------------- Hoja Coverage
    def _sheet_coverage(self, wb, cr) -> None:
        ws = wb.create_sheet("Coverage")

        headers = ["Metrica", "Valor", "Umbral", "Estado"]
        widths  = [38, 12, 12, 12]

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = HEADER_ALIGN
            cell.border = THIN_BORDER
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        if cr is None:
            return

        status_fills = {
            "OK":    PatternFill("solid", fgColor="44BB44"),
            "WARN":  PatternFill("solid", fgColor="FFAA00"),
            "FALTA": PatternFill("solid", fgColor="FF4444"),
        }
        status_fonts = {
            "OK":    Font(color="FFFFFF", bold=True, size=9),
            "WARN":  Font(color="000000", bold=True, size=9),
            "FALTA": Font(color="FFFFFF", bold=True, size=9),
        }

        metrics = [
            ("Cobertura de Requerimientos",
             f"{cr.requirements_coverage_pct:.1f}%", "100%",
             "OK" if cr.requirements_coverage_pct >= 100 else "FALTA"),
            ("Cobertura BVA",
             f"{cr.bva_coverage_pct:.1f}%", "100%",
             "OK" if cr.bva_coverage_pct >= 100 else "WARN"),
            ("Cobertura Particiones de Equivalencia",
             f"{cr.partition_coverage_pct:.1f}%", "100%",
             "OK" if cr.partition_coverage_pct >= 100 else "WARN"),
            ("Cobertura Tabla de Decision",
             f"{cr.decision_table_coverage_pct:.1f}%", "100%",
             "OK" if cr.decision_table_coverage_pct >= 100 else "WARN"),
            ("Cobertura Risk-Based",
             f"{cr.risk_based_coverage_pct:.1f}%", "100%",
             "OK" if cr.risk_based_coverage_pct >= 100 else "WARN"),
            ("Reglas con gap (sin casos)",
             str(len(cr.gaps)), "0",
             "OK" if not cr.gaps else "FALTA"),
        ]

        for row_idx, (metric, value, threshold, status) in enumerate(metrics, 2):
            for col_idx, val in enumerate([metric, value, threshold, status], 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font      = DATA_FONT
                cell.alignment = DATA_ALIGN
                cell.border    = THIN_BORDER
                if col_idx == 4:
                    cell.fill = status_fills.get(status, PatternFill())
                    cell.font = status_fonts.get(status, DATA_FONT)

    # ---------------------------------------------------------------- CSV
    def generate_csv(self, xlsx_path: str, csv_path: str) -> int:
        """
        Genera el CSV EXCLUSIVAMENTE desde el Excel ya generado (nunca desde memoria).
        Solo exporta la hoja 'Test Cases'. Retorna número de filas de datos.
        """
        wb = openpyxl.load_workbook(xlsx_path, read_only=True)
        ws = wb["Test Cases"]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        with open(csv_path, "w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.writer(fh, quoting=csv.QUOTE_ALL)
            for row in rows:
                writer.writerow([str(v) if v is not None else "" for v in row])

        return max(0, len(rows) - 1)
